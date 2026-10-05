"""文档解析：把上传的文件转为纯文本。"""

from __future__ import annotations

import io
from pathlib import Path

#: 按纯文本读取的扩展名（文档 + 常见代码 / 配置文件）
TEXT_EXTENSIONS = {
    ".txt", ".md", ".markdown", ".rst", ".csv", ".tsv", ".json", ".jsonl", ".log",
    ".yaml", ".yml", ".toml", ".ini", ".cfg", ".conf", ".xml", ".sql",
    ".py", ".js", ".ts", ".tsx", ".jsx", ".java", ".go", ".rs", ".c", ".h", ".cpp",
    ".hpp", ".cs", ".rb", ".php", ".sh", ".ps1", ".bat", ".kt", ".swift", ".scala",
    ".vue", ".css", ".scss", ".less", ".dockerfile", ".gradle", ".proto",
}
HTML_EXTENSIONS = {".html", ".htm", ".xhtml"}
SUPPORTED_EXTENSIONS = TEXT_EXTENSIONS | HTML_EXTENSIONS | {".pdf", ".docx"}


class UnsupportedFileType(ValueError):
    """不支持的文件类型。"""


def decode_text(data: bytes) -> str:
    for enc in ("utf-8-sig", "gb18030"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def _pdf_text(data: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    pages = []
    for i, page in enumerate(reader.pages, 1):
        text = (page.extract_text() or "").strip()
        if text:
            pages.append(f"[第 {i} 页]\n{text}")
    return "\n\n".join(pages)


def _docx_text(data: bytes) -> str:
    import docx

    document = docx.Document(io.BytesIO(data))
    parts = [p.text for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            if any(cells):
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def extract_text(filename: str, data: bytes) -> str:
    """按扩展名解析文件内容。

    Raises:
        UnsupportedFileType: 扩展名不受支持
    """
    name = filename.lower()
    ext = Path(name).suffix
    if name.endswith("dockerfile"):
        ext = ".dockerfile"
    if ext == ".pdf":
        return _pdf_text(data)
    if ext == ".docx":
        return _docx_text(data)
    if ext in HTML_EXTENSIONS:
        from app.capabilities.web.fetch import html_to_text

        title, text = html_to_text(decode_text(data))
        return f"# {title}\n\n{text}" if title else text
    if ext in TEXT_EXTENSIONS:
        return decode_text(data)
    raise UnsupportedFileType(
        f"不支持的文件类型: {ext or '无扩展名'}（支持 txt/md/pdf/docx/html 及常见代码文件）"
    )
